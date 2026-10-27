"""Tests for the document intelligence pipeline."""

from pathlib import Path

import pytest
from docx import Document

from src.chunking import chunk_document, fixed_chunks, sentence_chunks
from src.evaluate import _norm, score_question
from src.loaders import load_corpus, load_document


def test_fixed_chunks_overlap_and_no_split_words():
    text = " ".join(["word"] * 200)  # 999 chars
    chunks = fixed_chunks(text, size=100, overlap=20)
    assert len(chunks) > 1
    for chunk in chunks:
        assert "wor" not in chunk.replace("word", "")
    assert all(len(c) <= 100 for c in chunks)


def test_sentence_chunks_never_split_sentences():
    text = "First sentence here. Second sentence here! Is this the third?"
    chunks = sentence_chunks(text, max_chars=25)
    assert chunks == [
        "First sentence here.",
        "Second sentence here!",
        "Is this the third?",
    ]


def test_loaders_txt_docx_pdf(tmp_path):
    (tmp_path / "a.txt").write_text("Title\n\nPara one here.\n\nPara two here.\n")
    doc = Document()
    doc.add_heading("Doc Title", level=1)
    doc.add_paragraph("A docx paragraph.")
    doc.save(str(tmp_path / "b.docx"))
    import pymupdf
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 523, 770), "A pdf page.")
    pdf.save(str(tmp_path / "c.pdf"))
    assert len(load_document(tmp_path / "a.txt")) == 3
    assert load_document(tmp_path / "b.docx")[1][0] == "A docx paragraph."
    assert "pdf page" in load_document(tmp_path / "c.pdf")[0][0]


def test_corpus_has_30_docs_in_three_formats():
    docs = load_corpus(Path(__file__).resolve().parents[1] / "corpus")
    assert len(docs) == 30
    suffixes = {Path(d["path"]).suffix for d in docs}
    assert suffixes == {".pdf", ".docx", ".txt"}
    assert all(d["units"] for d in docs)


def test_chunk_document_citations():
    doc = {"doc_id": "demo", "units": [("Hello world. This is a test.", "p. 1")]}
    chunks = chunk_document(doc, "sentence")
    assert chunks[0]["doc_id"] == "demo"
    assert chunks[0]["locator"] == "p. 1"


def test_score_question_hit_and_correct():
    result = {
        "answer": "Employees receive 25 days of annual leave per year.",
        "results": [
            {"doc_id": "other"},
            {"doc_id": "leave_policy"},
        ],
    }
    hit, top1, correct = score_question(
        result, "leave_policy", "25 days of annual leave"
    )
    assert hit and top1 and correct
    hit, top1, correct = score_question(result, "moon_landing", "20 July 1969")
    assert not hit and not top1 and not correct


def test_norm_ignores_case_and_punctuation():
    assert _norm("Hello, World!") == _norm("hello world")


def test_end_to_end_ask(tmp_path):
    """Full path on a tiny corpus: hybrid search + rerank + citation."""
    from src.chunking import chunk_corpus
    from src.qa import ask
    from src.retrieval import Index, Reranker

    (tmp_path / "cats.txt").write_text(
        "Cats\n\nCats sleep for about 16 hours per day.\n\nDogs are loyal animals."
    )
    docs = load_corpus(tmp_path)
    index = Index(chunk_corpus(docs, "sentence"))
    result = ask("How long do cats sleep per day?", index, Reranker())
    assert "16 hours" in result["answer"]
    assert result["source_doc"] == "cats"
